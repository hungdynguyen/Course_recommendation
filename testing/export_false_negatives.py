#!/usr/bin/env python3
"""
Script xuất toàn bộ mẫu False Negative từ training dataset ra file Excel.
False Negative: mẫu mà negative score >= positive score (noise trong data).

Usage:
    python testing/export_false_negatives.py
    python testing/export_false_negatives.py --dataset data/processed/training_dataset/train_dataset_80.json
    python testing/export_false_negatives.py --sample 5000  # chỉ chạy trên 5000 mẫu đầu
"""

import argparse
import json
import os
import sys
from pathlib import Path

# Force offline mode — model is already cached locally; prevent HF Hub network calls
os.environ.setdefault("HF_HUB_OFFLINE", "1")
os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")

import numpy as np
import pandas as pd
from sentence_transformers import SentenceTransformer
from tqdm import tqdm

# ============== DEFAULTS ==============
DEFAULT_DATASET = "data/processed/training_dataset/train_dataset_80.json"
DEFAULT_MODEL = "Qwen/Qwen3-Embedding-0.6B"
DEFAULT_OUTPUT_DIR = Path("logs/dataset_analysis")
BATCH_SIZE = 32   # small batch to avoid GPU OOM

# ============================================================
# CLI args
# ============================================================
parser = argparse.ArgumentParser(description="Export False Negative samples to Excel")
parser.add_argument("--dataset", default=DEFAULT_DATASET, help="Path to training dataset JSON")
parser.add_argument("--model", default=DEFAULT_MODEL, help="SentenceTransformer model name/path")
parser.add_argument("--output_dir", default=str(DEFAULT_OUTPUT_DIR), help="Output directory")
parser.add_argument("--sample", type=int, default=None,
                    help="Only process first N samples (for quick testing)")
parser.add_argument("--threshold", type=float, default=0.0,
                    help="Export samples where margin <= threshold (default 0.0 = strict FN only)")
parser.add_argument("--device", default=None,
                    help="Device: 'cuda', 'cpu', or None (auto). Use 'cpu' if GPU OOM.")
parser.add_argument("--batch_size", type=int, default=BATCH_SIZE,
                    help=f"Encoding batch size (default {BATCH_SIZE})")
args = parser.parse_args()

DATASET_PATH   = Path(args.dataset)
MODEL_NAME     = args.model
OUTPUT_DIR     = Path(args.output_dir)
SAMPLE_SIZE    = args.sample
MARGIN_THRESH  = args.threshold
DEVICE         = args.device
BATCH_SIZE     = args.batch_size

print("=" * 80)
print("🔍 EXPORT FALSE NEGATIVES → EXCEL")
print("=" * 80)
print(f"Dataset  : {DATASET_PATH}")
print(f"Model    : {MODEL_NAME}")
print(f"Threshold: margin <= {MARGIN_THRESH}")
print(f"Sample   : {'Full dataset' if SAMPLE_SIZE is None else SAMPLE_SIZE}")
print(f"Device   : {DEVICE or 'auto'} | Batch size: {BATCH_SIZE}")
print()

# ============================================================
# 1. Load dataset
# ============================================================
print("📂 Loading dataset...")
if not DATASET_PATH.exists():
    print(f"❌ File not found: {DATASET_PATH}", file=sys.stderr)
    sys.exit(1)

with open(DATASET_PATH, "r", encoding="utf-8") as f:
    data = json.load(f)

if SAMPLE_SIZE:
    data = data[:SAMPLE_SIZE]

print(f"✓ Loaded {len(data):,} samples\n")

# ============================================================
# 2. Load model
# ============================================================
import torch

print(f"🤖 Loading model: {MODEL_NAME}...")
model = SentenceTransformer(MODEL_NAME, trust_remote_code=True, device=DEVICE)
actual_device = str(next(model.parameters()).device)
print(f"✓ Model loaded on {actual_device}\n")

# ============================================================
# 3. Encode & compute similarity scores
# ============================================================
def encode_to_numpy(texts, label):
    """Encode texts → L2-normalized float32 numpy array, free GPU cache after."""
    print(f"  - Encoding {label}...")
    emb = model.encode(
        texts,
        convert_to_numpy=True,
        normalize_embeddings=True,
        batch_size=BATCH_SIZE,
        show_progress_bar=True,
    )
    # Free GPU memory immediately before next encode call
    if torch.cuda.is_available():
        torch.cuda.empty_cache()
    return emb.astype(np.float32)

print("🧮 Computing similarity scores...")
queries   = [item["query"]    for item in data]
positives = [item["positive"] for item in data]
negatives = [item["negative"] for item in data]

q_emb = encode_to_numpy(queries,   "queries")
p_emb = encode_to_numpy(positives, "positives")
n_emb = encode_to_numpy(negatives, "negatives")

# Cosine similarity via dot product (embeddings are L2-normalized)
print("  - Computing cosine similarities (CPU numpy)...")
pos_scores = np.einsum('ij,ij->i', q_emb, p_emb).astype(np.float32)
neg_scores = np.einsum('ij,ij->i', q_emb, n_emb).astype(np.float32)
margins    = pos_scores - neg_scores

# Free embedding arrays
del q_emb, p_emb, n_emb

print("✓ Done\n")

# ============================================================
# 4. Identify false negatives
# ============================================================
fn_mask    = margins <= MARGIN_THRESH
fn_indices = np.where(fn_mask)[0]
fn_count   = len(fn_indices)
fn_ratio   = fn_count / len(data) * 100

print(f"⚠️  False Negatives (margin ≤ {MARGIN_THRESH}): {fn_count:,} / {len(data):,}  ({fn_ratio:.2f}%)\n")

if fn_count == 0:
    print("✅ Không có mẫu nào bị False Negative. Không cần xuất file.")
    sys.exit(0)

# ============================================================
# 5. Build DataFrame
# ============================================================
print("📋 Building DataFrame...")
rows = []
for idx in tqdm(fn_indices, desc="Building rows"):
    item = data[idx]
    meta = item.get("metadata", {})
    rows.append({
        "index"         : int(idx),
        "margin"        : round(float(margins[idx]), 6),
        "pos_score"     : round(float(pos_scores[idx]), 6),
        "neg_score"     : round(float(neg_scores[idx]), 6),
        "query"         : item.get("query", ""),
        "positive"      : item.get("positive", ""),
        "negative"      : item.get("negative", ""),
        "skill_name"    : meta.get("skill_name", ""),
        "skill_type"    : meta.get("skill_type", ""),
        "skill_uri"     : meta.get("skill_uri", ""),
        "generated_at"  : meta.get("generated_at", ""),
        "model"         : meta.get("model", ""),
    })

df = pd.DataFrame(rows)
# Sort by margin ascending (worst first)
df.sort_values("margin", ascending=True, inplace=True)
df.reset_index(drop=True, inplace=True)

# ============================================================
# 6. Save to Excel
# ============================================================
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
thresh_tag = f"thresh{MARGIN_THRESH:.2f}".replace(".", "p")
output_path = OUTPUT_DIR / f"false_negatives_{thresh_tag}.xlsx"

print(f"\n💾 Saving to {output_path}...")

with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
    # ---- Sheet 1: False Negative samples ----
    df.to_excel(writer, sheet_name="False Negatives", index=False)

    ws = writer.sheets["False Negatives"]

    # Column widths
    col_widths = {
        "A": 8,   # index
        "B": 10,  # margin
        "C": 10,  # pos_score
        "D": 10,  # neg_score
        "E": 55,  # query
        "F": 55,  # positive
        "G": 55,  # negative
        "H": 30,  # skill_name
        "I": 20,  # skill_type
        "J": 55,  # skill_uri
        "K": 22,  # generated_at
        "L": 20,  # model
    }
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    # Wrap text + vertical align top for text columns
    from openpyxl.styles import Alignment, PatternFill, Font
    wrap_align  = Alignment(wrap_text=True, vertical="top")
    center_align = Alignment(horizontal="center", vertical="top")

    # Header styling
    header_fill = PatternFill("solid", fgColor="667EEA")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Freeze header row
    ws.freeze_panes = "A2"

    # Row-level coloring & alignment
    red_fill    = PatternFill("solid", fgColor="FFCDD2")   # very negative margin
    orange_fill = PatternFill("solid", fgColor="FFE0B2")   # slightly negative
    normal_fill = PatternFill("solid", fgColor="FFFFFF")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        margin_val = row[1].value  # column B = margin
        if margin_val is not None:
            if margin_val < -0.05:
                fill = red_fill
            elif margin_val < 0:
                fill = orange_fill
            else:
                fill = normal_fill
        else:
            fill = normal_fill

        for i, cell in enumerate(row):
            cell.fill = fill
            if i in (4, 5, 6):  # query, positive, negative
                cell.alignment = wrap_align
            else:
                cell.alignment = Alignment(vertical="top")

    # ---- Sheet 2: Summary ----
    summary_data = {
        "Metric": [
            "Dataset path",
            "Model",
            "Total samples",
            "False Negative count",
            "False Negative ratio (%)",
            "Threshold (margin ≤)",
            "Worst margin",
            "Mean margin (FN only)",
            "Median margin (FN only)",
        ],
        "Value": [
            str(DATASET_PATH),
            MODEL_NAME,
            len(data),
            fn_count,
            round(fn_ratio, 4),
            MARGIN_THRESH,
            round(float(margins[fn_indices].min()), 6),
            round(float(margins[fn_indices].mean()), 6),
            round(float(np.median(margins[fn_indices])), 6),
        ],
    }
    df_summary = pd.DataFrame(summary_data)
    df_summary.to_excel(writer, sheet_name="Summary", index=False)

    ws2 = writer.sheets["Summary"]
    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 55
    for cell in ws2[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

print(f"✓ Saved {fn_count:,} false negative samples → {output_path}")
print()
print("=" * 80)
print("✅ HOÀN TẤT!")
print("=" * 80)
